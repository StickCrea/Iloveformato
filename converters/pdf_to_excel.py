import re

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Font

MONEY_RE = re.compile(r"-?\d{0,3}(?:,\d{3})*\.\d{2}")
DATE_LINE_RE = re.compile(r"^(\d{1,2})/(\d{2})\s+(.*)$")
PERIOD_RE = re.compile(
    r"DESDE:\s*(\d{4})/(\d{2})/\d{2}\s+HASTA:\s*(\d{4})/(\d{2})/\d{2}"
)

STATEMENT_HEADERS = ["Fecha", "Descripcion", "Valor", "Saldo"]
MIN_STATEMENT_ROWS = 5


def _parse_money(text):
    return float(text.replace(",", ""))


def _resolve_year(month, period):
    if period is None:
        return None
    desde_year, desde_month, hasta_year, hasta_month = period
    if desde_year == hasta_year:
        return desde_year
    return desde_year if month >= desde_month else hasta_year


def _extract_statement_rows(input_path):
    """Parse bank-statement-style PDFs (Fecha/Descripcion/.../Valor/Saldo lines)."""
    rows = []
    period = None
    with pdfplumber.open(input_path) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            for line in text.splitlines():
                period_match = PERIOD_RE.search(line)
                if period_match:
                    dy, dm, hy, hm = period_match.groups()
                    period = (int(dy), int(dm), int(hy), int(hm))
                    continue

                date_match = DATE_LINE_RE.match(line)
                if not date_match:
                    continue

                day, month, resto = date_match.groups()
                monies = MONEY_RE.findall(resto)
                if len(monies) < 2:
                    continue

                valor_txt, saldo_txt = monies[-2], monies[-1]
                idx_saldo = resto.rfind(saldo_txt)
                idx_valor = resto[:idx_saldo].rfind(valor_txt)
                descripcion = resto[:idx_valor].strip()
                if not descripcion:
                    continue

                year = _resolve_year(int(month), period)
                fecha = f"{day}/{month}/{year}" if year else f"{day}/{month}"

                rows.append(
                    (fecha, descripcion, _parse_money(valor_txt), _parse_money(saldo_txt))
                )
    return rows


def _write_statement_sheet(wb, rows):
    sheet = wb.create_sheet(title="Movimientos")
    sheet.append(STATEMENT_HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    for column, width in zip("ABCD", (12, 55, 16, 16)):
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2, min_col=3, max_col=4):
        for cell in row:
            cell.number_format = "#,##0.00"
    sheet.freeze_panes = "A2"


def _write_generic_sheets(wb, input_path):
    with pdfplumber.open(input_path) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            tables = page.extract_tables()
            if tables:
                for table_index, table in enumerate(tables, start=1):
                    sheet_name = f"Pagina{page_number}_Tabla{table_index}"[:31]
                    sheet = wb.create_sheet(title=sheet_name)
                    for row in table:
                        sheet.append(["" if cell is None else cell for cell in row])
            else:
                text = page.extract_text() or ""
                sheet = wb.create_sheet(title=f"Pagina{page_number}"[:31])
                for line in text.splitlines():
                    sheet.append([line])


def convert(input_path: str, output_path: str) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    statement_rows = _extract_statement_rows(input_path)
    if len(statement_rows) >= MIN_STATEMENT_ROWS:
        _write_statement_sheet(wb, statement_rows)
    else:
        _write_generic_sheets(wb, input_path)

    if not wb.sheetnames:
        wb.create_sheet(title="Sin contenido")

    wb.save(output_path)
