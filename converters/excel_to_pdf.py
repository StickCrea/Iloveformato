from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, PageBreak

MAX_COLS_PER_CHUNK = 10


def _chunk_columns(rows, chunk_size):
    if not rows:
        return [rows]
    total_cols = max(len(row) for row in rows)
    chunks = []
    for start in range(0, total_cols, chunk_size):
        end = start + chunk_size
        chunks.append([row[start:end] for row in rows])
    return chunks or [rows]


def convert(input_path: str, output_path: str) -> None:
    wb = load_workbook(input_path, data_only=True)
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(
        output_path,
        pagesize=landscape(A4),
        leftMargin=1.5 * cm,
        rightMargin=1.5 * cm,
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm,
    )

    story = []
    for sheet_index, sheet in enumerate(wb.worksheets):
        rows = [
            ["" if cell is None else str(cell) for cell in row]
            for row in sheet.iter_rows(values_only=True)
        ]
        rows = [row for row in rows if any(cell != "" for cell in row)]

        if sheet_index > 0:
            story.append(PageBreak())
        story.append(Paragraph(sheet.title, styles["Heading2"]))

        if not rows:
            story.append(Paragraph("(hoja vacia)", styles["Normal"]))
            continue

        for chunk in _chunk_columns(rows, MAX_COLS_PER_CHUNK):
            table = Table(chunk, repeatRows=1)
            table.setStyle(
                TableStyle(
                    [
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2f6feb")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("FONTSIZE", (0, 0), (-1, -1), 7),
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ]
                )
            )
            story.append(table)

    doc.build(story)
